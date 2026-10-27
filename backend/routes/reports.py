import json
import re
from datetime import date, datetime
from io import BytesIO
from xml.sax.saxutils import escape

from flask import Blueprint, current_app, request, send_file, session
from mysql.connector import Error
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from backend.db import get_connection
from backend.routes.tasks import VALID_DUE_FILTERS, VALID_PRIORITIES, VALID_STATUSES
from backend.security import login_required

reports_bp = Blueprint("reports", __name__)


def _filters():
    filters = {
        "search": request.args.get("search", "").strip(),
        "status": request.args.get("status", "").strip(),
        "priority": request.args.get("priority", "").strip(),
        "due_date": request.args.get("due_date", "").strip(),
        "sort": request.args.get("sort", "created_at").strip(),
    }
    if filters["status"] and filters["status"] not in VALID_STATUSES:
        return None, "Invalid status filter."
    if filters["priority"] and filters["priority"] not in VALID_PRIORITIES:
        return None, "Invalid priority filter."
    if filters["due_date"] and filters["due_date"] not in VALID_DUE_FILTERS:
        return None, "Invalid due date filter."
    return filters, None


def _load_tasks(user_id, filters):
    sort_columns = {"due_date": "due_date", "created_at": "created_at", "priority": "priority"}
    order_column = sort_columns.get(filters["sort"], "created_at")
    order = "FIELD(priority, 'high', 'medium', 'low'), created_at DESC" if filters["sort"] == "priority" else f"{order_column} ASC"
    clauses, values = ["user_id = %s"], [user_id]
    if filters["search"]:
        clauses.append("title LIKE %s")
        values.append(f"%{filters['search']}%")
    if filters["status"]:
        clauses.append("status = %s")
        values.append(filters["status"])
    if filters["priority"]:
        clauses.append("priority = %s")
        values.append(filters["priority"])
    if filters["due_date"] == "overdue":
        clauses.append("due_date < CURRENT_DATE AND status <> 'completed'")
    elif filters["due_date"] == "today":
        clauses.append("due_date = CURRENT_DATE")
    elif filters["due_date"] == "upcoming":
        clauses.append("due_date >= CURRENT_DATE AND due_date <= DATE_ADD(CURRENT_DATE, INTERVAL 7 DAY)")
    elif filters["due_date"] == "no_due":
        clauses.append("due_date IS NULL")
    sql = f"""
        SELECT id, title, description, status, priority, due_date, created_at, updated_at
        FROM tasks
        WHERE {' AND '.join(clauses)}
        ORDER BY {order}
    """
    with get_connection() as connection:
        cursor = connection.cursor(dictionary=True)
        cursor.execute(sql, tuple(values))
        tasks = cursor.fetchall()
        cursor.close()
    return tasks


def _format_date(value):
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    return str(value) if value else ""


def _safe_filename(prefix, extension):
    return f"taskflow-{prefix}-{date.today().isoformat()}.{extension}"


def _excel_file(tasks):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Task Report"
    sheet.merge_cells("A1:I1")
    sheet["A1"] = "TaskFlow Task Export"
    sheet["A1"].font = Font(size=16, bold=True, color="FFFFFF")
    sheet["A1"].fill = PatternFill("solid", fgColor="7556DF")
    sheet["A1"].alignment = Alignment(horizontal="center")
    sheet.merge_cells("A2:I2")
    sheet["A2"] = f"Generated {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    sheet["A2"].font = Font(italic=True, color="666666")

    headers = [
        "Task ID", "Title", "Description", "Priority", "Status",
        "Due Date", "Created Date", "Updated Date", "AI Subtasks",
    ]
    header_row = 4
    for column, header in enumerate(headers, 1):
        cell = sheet.cell(header_row, column, header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="5B43B8")
        cell.alignment = Alignment(horizontal="center")
    for row_number, task in enumerate(tasks, header_row + 1):
        values = [
            task["id"], task["title"], task["description"] or "", task["priority"],
            task["status"], task["due_date"], task["created_at"], task["updated_at"], "",
        ]
        for column, value in enumerate(values, 1):
            cell = sheet.cell(row_number, column, value)
            cell.alignment = Alignment(vertical="top", wrap_text=column in (2, 3, 9))
        for column in (6, 7, 8):
            sheet.cell(row_number, column).number_format = "yyyy-mm-dd"
    widths = [10, 28, 52, 14, 16, 15, 20, 20, 30]
    for column, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = "A5"
    sheet.auto_filter.ref = f"A{header_row}:I{max(header_row, header_row + len(tasks))}"
    sheet.row_dimensions[1].height = 25
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def _verified_metrics(tasks):
    today = date.today()
    total = len(tasks)
    completed = sum(task["status"] == "completed" for task in tasks)
    pending = total - completed
    overdue = sum(
        task["status"] != "completed"
        and task["due_date"] is not None
        and task["due_date"] < today
        for task in tasks
    )
    return {
        "total": total,
        "completed": completed,
        "pending": pending,
        "overdue": overdue,
        "completion_rate": round(completed / total * 100) if total else 0,
    }


def _ai_insights(tasks, metrics):
    if not current_app.config.get("GEMINI_API_KEY"):
        return None, "Gemini AI is not configured; this report contains factual metrics only."
    task_context = [
        {
            "title": task["title"][:160],
            "status": task["status"],
            "priority": task["priority"],
            "due_date": _format_date(task["due_date"]) or None,
        }
        for task in tasks
    ]
    prompt = (
        "Return only valid JSON with exactly these string/list keys: executive_summary (string), "
        "progress_highlights (array of strings), potential_blockers (array of strings), "
        "recommended_priorities (array of strings), actionable_next_steps (array of strings). "
        "Analyze the task context, but do not calculate, alter, or state task counts, completion rates, "
        "or dates. Do not invent tasks or dates. Keep each list to at most 5 concise items. "
        f"Verified metrics are supplied separately and must not be repeated: {json.dumps(metrics)}. "
        f"Task context: {json.dumps(task_context)}"
    )
    try:
        from google import genai
        from google.genai import types

        client = genai.Client(
            api_key=current_app.config["GEMINI_API_KEY"],
            http_options=types.HttpOptions(timeout=10000),
        )
        response = client.models.generate_content(
            model=current_app.config["GEMINI_MODEL"],
            contents=prompt,
            config=types.GenerateContentConfig(response_mime_type="application/json", temperature=0.2),
        )
        data = json.loads(response.text)
        required = ("executive_summary", "progress_highlights", "potential_blockers", "recommended_priorities", "actionable_next_steps")
        if not isinstance(data, dict) or not isinstance(data.get("executive_summary"), str):
            raise ValueError("Invalid report structure")
        for key in required[1:]:
            if not isinstance(data.get(key), list):
                raise ValueError("Invalid report list")
            if not all(isinstance(item, str) and item.strip() for item in data[key]):
                raise ValueError("Invalid report list")
            data[key] = [item.strip()[:500] for item in data[key][:5]]
        data["executive_summary"] = data["executive_summary"].strip()[:1500]
        ai_text = " ".join([data["executive_summary"], *sum((data[key] for key in required[1:]), [])])
        if re.search(r"\b\d+(?:\.\d+)?%?\b|\b\d{4}-\d{2}-\d{2}\b", ai_text):
            raise ValueError("AI report contains unverified numeric or date claims")
        return data, None
    except (json.JSONDecodeError, TypeError, ValueError):
        current_app.logger.warning("AI provider returned an invalid productivity report")
        return None, "Gemini returned an invalid analysis; this report contains factual metrics only."
    except Exception:
        current_app.logger.warning("AI productivity report request failed", exc_info=True)
        return None, "Gemini was unavailable; this report contains factual metrics only."


def _pdf_file(tasks, metrics, insights, limitation):
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="ReportTitle", parent=styles["Title"], alignment=TA_CENTER, textColor=colors.HexColor("#5B43B8"), spaceAfter=8))
    styles.add(ParagraphStyle(name="Small", parent=styles["BodyText"], fontSize=8, leading=10, textColor=colors.HexColor("#555555")))
    styles.add(ParagraphStyle(name="Section", parent=styles["Heading2"], textColor=colors.HexColor("#5B43B8"), spaceBefore=12, spaceAfter=6))
    styles.add(ParagraphStyle(name="Cell", parent=styles["BodyText"], fontSize=8, leading=10))
    story = [
        Paragraph("TaskFlow Productivity Report", styles["ReportTitle"]),
        Paragraph(f"Generated {date.today().strftime('%B %d, %Y')}", styles["Small"]),
        Spacer(1, 8),
    ]
    kpi_data = [
        ["Total tasks", "Completed", "Pending", "Overdue", "Completion rate"],
        [str(metrics["total"]), str(metrics["completed"]), str(metrics["pending"]), str(metrics["overdue"]), f"{metrics['completion_rate']}%"],
    ]
    kpi_table = Table(kpi_data, colWidths=[35 * mm] * 5)
    kpi_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#5B43B8")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#D8D0F5")),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#D8D0F5")),
        ("BACKGROUND", (0, 1), (-1, 1), colors.HexColor("#F3F0FF")),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
        ("TOPPADDING", (0, 0), (-1, -1), 7),
    ]))
    story.extend([kpi_table, Spacer(1, 8), Paragraph("Status breakdown", styles["Section"])])
    status_data = [["Status", "Tasks"], ["To do", str(sum(task["status"] == "todo" for task in tasks))], ["In progress", str(sum(task["status"] == "in_progress" for task in tasks))], ["Completed", str(metrics["completed"])]]
    story.append(Table(status_data, colWidths=[80 * mm, 40 * mm], style=TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EEE9FF")), ("GRID", (0, 0), (-1, -1), 0.25, colors.grey), ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold")])))
    story.append(Paragraph("AI insights", styles["Section"]))
    if insights:
        story.append(Paragraph(escape(insights["executive_summary"]), styles["BodyText"]))
        for heading, key in (("Progress highlights", "progress_highlights"), ("Potential blockers", "potential_blockers"), ("Recommended priorities", "recommended_priorities"), ("Actionable next steps", "actionable_next_steps")):
            story.append(Paragraph(heading, styles["Heading3"]))
            story.extend(Paragraph(f"- {escape(item)}", styles["BodyText"]) for item in insights[key])
    else:
        story.append(Paragraph(escape(limitation or "No AI insights were generated."), styles["BodyText"]))
    story.extend([Paragraph("Task summary", styles["Section"])])
    task_rows = [["ID", "Title", "Priority", "Status", "Due date"]]
    task_rows.extend([[str(task["id"]), Paragraph(escape(task["title"]), styles["Cell"]), task["priority"], task["status"], _format_date(task["due_date"]) or "—"] for task in tasks])
    task_table = Table(task_rows, colWidths=[12 * mm, 82 * mm, 25 * mm, 28 * mm, 28 * mm], repeatRows=1)
    task_table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#5B43B8")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("GRID", (0, 0), (-1, -1), 0.25, colors.grey), ("VALIGN", (0, 0), (-1, -1), "TOP"), ("FONTSIZE", (0, 0), (-1, -1), 8)]))
    story.append(task_table)
    output = BytesIO()
    document = SimpleDocTemplate(output, pagesize=A4, rightMargin=15 * mm, leftMargin=15 * mm, topMargin=15 * mm, bottomMargin=15 * mm, title="TaskFlow Productivity Report")
    document.build(story)
    output.seek(0)
    return output


@reports_bp.get("/reports/tasks.xlsx")
@login_required
def export_excel():
    filters, error = _filters()
    if error:
        return {"success": False, "error": {"message": error}}, 400
    try:
        tasks = _load_tasks(session["user_id"], filters)
        return send_file(_excel_file(tasks), as_attachment=True, download_name=_safe_filename("tasks", "xlsx"), mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    except Error:
        current_app.logger.exception("Unable to export tasks to Excel")
        return {"success": False, "error": {"message": "Database unavailable."}}, 503
    except (OSError, ValueError):
        current_app.logger.exception("Unable to generate Excel export")
        return {"success": False, "error": {"message": "Unable to generate the Excel export."}}, 500


@reports_bp.get("/reports/productivity.pdf")
@login_required
def productivity_pdf():
    filters, error = _filters()
    if error:
        return {"success": False, "error": {"message": error}}, 400
    try:
        tasks = _load_tasks(session["user_id"], filters)
        metrics = _verified_metrics(tasks)
        insights, limitation = _ai_insights(tasks, metrics)
        output = _pdf_file(tasks, metrics, insights, limitation)
        response = send_file(output, as_attachment=True, download_name=_safe_filename("productivity-report", "pdf"), mimetype="application/pdf")
        if limitation:
            response.headers["X-TaskFlow-Report-Note"] = limitation
        return response
    except Error:
        current_app.logger.exception("Unable to generate productivity report")
        return {"success": False, "error": {"message": "Database unavailable."}}, 503
    except (OSError, ValueError):
        current_app.logger.exception("Unable to generate PDF report")
        return {"success": False, "error": {"message": "Unable to generate the PDF report."}}, 500
