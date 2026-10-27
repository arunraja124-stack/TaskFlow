import uuid
import threading
from datetime import date, datetime, timedelta, timezone
from io import BytesIO

import mysql.connector
from openpyxl import load_workbook
import pytest
from zoneinfo import ZoneInfo

from app import create_app
from backend.attendance import business_date, display_time
from backend.routes.reports import _verified_metrics
from config import Config


@pytest.fixture()
def client(monkeypatch):
    database = f"taskflow_test_{uuid.uuid4().hex[:12]}"
    admin_config = dict(Config.DB_CONFIG)
    admin_config.pop("database", None)
    connection = mysql.connector.connect(**admin_config)
    cursor = connection.cursor()
    cursor.execute(f"CREATE DATABASE `{database}`")
    cursor.execute(f"USE `{database}`")
    cursor.execute("""CREATE TABLE users (
        id INT UNSIGNED AUTO_INCREMENT PRIMARY KEY, name VARCHAR(100) NOT NULL,
        email VARCHAR(254) NOT NULL UNIQUE, password_hash VARCHAR(255) NOT NULL,
        role ENUM('member', 'admin') NOT NULL DEFAULT 'member',
        created_at TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP
    )""")
    cursor.execute("""CREATE TABLE tasks (
        id INT UNSIGNED AUTO_INCREMENT PRIMARY KEY, user_id INT UNSIGNED NOT NULL,
        title VARCHAR(160) NOT NULL, description VARCHAR(2000) NOT NULL DEFAULT '',
        status ENUM('todo', 'in_progress', 'completed') NOT NULL DEFAULT 'todo',
        priority ENUM('low', 'medium', 'high') NOT NULL DEFAULT 'medium',
        due_date DATE NULL, created_at TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP,
        updated_at TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP
    )""")
    cursor.execute("""CREATE TABLE task_activity (
        id BIGINT UNSIGNED AUTO_INCREMENT PRIMARY KEY, user_id INT UNSIGNED NOT NULL,
        task_id INT UNSIGNED NULL, action VARCHAR(40) NOT NULL,
        details VARCHAR(500) NOT NULL DEFAULT '',
        created_at TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP
    )""")
    cursor.execute("""CREATE TABLE notifications (
        id BIGINT UNSIGNED AUTO_INCREMENT PRIMARY KEY, user_id INT UNSIGNED NOT NULL,
        task_id INT UNSIGNED NULL, kind VARCHAR(40) NOT NULL,
        message VARCHAR(255) NOT NULL, is_read BOOLEAN NOT NULL DEFAULT FALSE,
        created_at TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP
    )""")
    cursor.execute("""CREATE TABLE user_preferences (
        user_id INT UNSIGNED PRIMARY KEY, deadline_reminders BOOLEAN NOT NULL DEFAULT TRUE,
        overdue_alerts BOOLEAN NOT NULL DEFAULT TRUE, default_view VARCHAR(20) NOT NULL DEFAULT 'dashboard',
        default_priority ENUM('low', 'medium', 'high') NOT NULL DEFAULT 'medium'
    )""")
    cursor.execute("""CREATE TABLE attendance (
        id BIGINT UNSIGNED AUTO_INCREMENT PRIMARY KEY, user_id INT UNSIGNED NOT NULL,
        attendance_date DATE NOT NULL, check_in_at DATETIME NOT NULL,
        check_out_at DATETIME NULL, created_at TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP,
        updated_at TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP,
        UNIQUE KEY uq_attendance_user_date (user_id, attendance_date),
        INDEX idx_attendance_date (attendance_date),
        INDEX idx_attendance_user_date (user_id, attendance_date),
        CONSTRAINT fk_attendance_user FOREIGN KEY (user_id) REFERENCES users(id) ON DELETE CASCADE
    )""")
    connection.commit()
    cursor.close()
    connection.close()
    test_config = dict(Config.DB_CONFIG)
    test_config["database"] = database
    monkeypatch.setattr(Config, "DB_CONFIG", test_config)
    app = create_app({"TESTING": True, "FRONTEND_ORIGIN": "*", "GEMINI_API_KEY": ""})
    yield app.test_client()
    cleanup = mysql.connector.connect(**admin_config)
    cleanup_cursor = cleanup.cursor()
    cleanup_cursor.execute(f"DROP DATABASE `{database}`")
    cleanup.commit()
    cleanup_cursor.close()
    cleanup.close()


def csrf(client):
    return client.get("/api/auth/csrf").json["data"]["csrf_token"]


def register(client, email="one@example.com"):
    token = csrf(client)
    result = client.post(
        "/api/auth/register",
        json={"name": "User One", "email": email, "password": "secure-pass-123", "password_confirmation": "secure-pass-123"},
        headers={"X-CSRF-Token": token},
    )
    assert result.status_code == 201
    return result


def test_health_and_anonymous_protection(client):
    assert client.get("/api/health").status_code == 200
    assert client.get("/api/tasks").status_code == 401


def test_registration_hashes_password_and_duplicate_is_generic(client):
    register(client)
    connection = mysql.connector.connect(**Config.DB_CONFIG)
    cursor = connection.cursor()
    cursor.execute("SELECT password_hash FROM users WHERE email = %s", ("one@example.com",))
    password_hash = cursor.fetchone()[0]
    assert password_hash != "secure-pass-123"
    cursor.close()
    connection.close()
    token = csrf(client)
    duplicate = client.post("/api/auth/register", json={"name": "Other", "email": "one@example.com", "password": "secure-pass-123", "password_confirmation": "secure-pass-123"}, headers={"X-CSRF-Token": token})
    assert duplicate.status_code == 409


def test_login_logout_and_crud(client):
    register(client)
    token = csrf(client)
    created = client.post("/api/tasks", json={"title": "Owned task", "priority": "high"}, headers={"X-CSRF-Token": token})
    assert created.status_code == 201
    task_id = created.json["data"]["id"]
    assert client.get(f"/api/tasks/{task_id}").status_code == 200
    token = csrf(client)
    assert client.put(f"/api/tasks/{task_id}", json={"title": "Updated task"}, headers={"X-CSRF-Token": token}).status_code == 200
    token = csrf(client)
    assert client.patch(f"/api/tasks/{task_id}/status", json={"status": "completed"}, headers={"X-CSRF-Token": token}).status_code == 200
    assert client.get("/api/analytics").json["data"]["completed"] == 1
    token = csrf(client)
    assert client.delete(f"/api/tasks/{task_id}", headers={"X-CSRF-Token": token}).status_code == 200
    token = csrf(client)
    assert client.post("/api/auth/logout", headers={"X-CSRF-Token": token}).status_code == 200
    assert client.get("/api/tasks").status_code == 401


def test_cross_user_task_is_not_accessible(client):
    register(client)
    token = csrf(client)
    task_id = client.post("/api/tasks", json={"title": "Private task"}, headers={"X-CSRF-Token": token}).json["data"]["id"]
    token = csrf(client)
    client.post("/api/auth/logout", headers={"X-CSRF-Token": token})
    token = csrf(client)
    register(client, "two@example.com")
    assert client.get(f"/api/tasks/{task_id}").status_code == 404
    token = csrf(client)
    assert client.patch(f"/api/tasks/{task_id}/status", json={"status": "completed"}, headers={"X-CSRF-Token": token}).status_code == 404


def test_validation_csrf_and_ai_configuration(client):
    register(client)
    assert client.post("/api/tasks", json={"title": "Missing token"}).status_code == 403
    assert client.post("/api/ai/suggest-tasks", json={"description": "Plan a release"}).status_code == 403
    token = csrf(client)
    assert client.post("/api/tasks", json={"description": "Missing title"}, headers={"X-CSRF-Token": token}).status_code == 400
    token = csrf(client)
    ai = client.post("/api/ai/suggest-tasks", json={"description": "Plan a release"}, headers={"X-CSRF-Token": token})
    assert ai.status_code == 503


def test_activity_calendar_and_reminders_are_user_scoped(client):
    register(client)
    token = csrf(client)
    created = client.post(
        "/api/tasks",
        json={"title": "Dated task", "due_date": (date.today() + timedelta(days=3)).isoformat()},
        headers={"X-CSRF-Token": token},
    )
    assert created.status_code == 201
    assert client.get("/api/activity").json["data"]["activity"][0]["action"] == "created"
    assert client.get("/api/calendar").json["data"]["tasks"][0]["title"] == "Dated task"
    assert client.get("/api/reminders").json["data"]["reminders"][0]["kind"] == "upcoming"
    assert client.get("/api/tasks?due_date=upcoming").json["data"]["count"] == 1
    assert client.get("/api/tasks?due_date=no_due").json["data"]["count"] == 0
    assert client.get("/api/tasks?due_date=invalid").status_code == 400


def test_task_filters_and_validation(client):
    register(client)
    token = csrf(client)
    first = client.post(
        "/api/tasks",
        json={"title": "High active", "priority": "high", "due_date": (date.today() + timedelta(days=1)).isoformat()},
        headers={"X-CSRF-Token": token},
    )
    assert first.status_code == 201
    token = csrf(client)
    second = client.post(
        "/api/tasks",
        json={"title": "Medium completed", "priority": "medium", "due_date": date.today().isoformat()},
        headers={"X-CSRF-Token": token},
    )
    assert second.status_code == 201
    token = csrf(client)
    assert client.patch(
        f"/api/tasks/{second.json['data']['id']}/status",
        json={"status": "completed"},
        headers={"X-CSRF-Token": token},
    ).status_code == 200
    assert client.get("/api/tasks?priority=high").json["data"]["count"] == 1
    assert client.get("/api/tasks?status=completed").json["data"]["count"] == 1
    assert client.get("/api/tasks?search=High").json["data"]["count"] == 1
    assert client.get("/api/tasks?priority=invalid").status_code == 400
    assert client.get("/api/tasks?status=invalid").status_code == 400


def test_authenticated_report_downloads_and_empty_lists(client):
    assert client.get("/api/reports/tasks.xlsx").status_code == 401
    assert client.get("/api/reports/productivity.pdf").status_code == 401
    register(client)

    excel = client.get("/api/reports/tasks.xlsx")
    assert excel.status_code == 200
    assert excel.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    workbook = load_workbook(BytesIO(excel.data))
    sheet = workbook["Task Report"]
    assert sheet["A1"].value == "TaskFlow Task Export"
    assert sheet["A5"].value is None
    assert sheet.freeze_panes == "A5"

    pdf = client.get("/api/reports/productivity.pdf")
    assert pdf.status_code == 200
    assert pdf.mimetype == "application/pdf"
    assert pdf.data.startswith(b"%PDF-")
    assert "Gemini AI is not configured" in pdf.headers["X-TaskFlow-Report-Note"]


def test_reports_use_real_user_tasks_and_support_special_content(client):
    register(client)
    token = csrf(client)
    created = client.post(
        "/api/tasks",
        json={"title": "Quarterly <review> & planning", "description": "Long " * 200, "priority": "high"},
        headers={"X-CSRF-Token": token},
    )
    assert created.status_code == 201

    excel = client.get("/api/reports/tasks.xlsx")
    sheet = load_workbook(BytesIO(excel.data), read_only=True)["Task Report"]
    rows = list(sheet.iter_rows(min_row=5, values_only=True))
    assert len(rows) == 1
    assert rows[0][1] == "Quarterly <review> & planning"
    assert rows[0][8] in ("", None)

    pdf = client.get("/api/reports/productivity.pdf")
    assert pdf.status_code == 200
    assert pdf.data.startswith(b"%PDF-")


def test_report_metrics_and_excel_filters_are_factual(client):
    register(client)
    token = csrf(client)
    created = client.post(
        "/api/tasks",
        json={"title": "High report task", "priority": "high"},
        headers={"X-CSRF-Token": token},
    )
    assert created.status_code == 201
    token = csrf(client)
    completed = client.post(
        "/api/tasks",
        json={"title": "Completed report task", "priority": "medium"},
        headers={"X-CSRF-Token": token},
    )
    assert completed.status_code == 201
    token = csrf(client)
    assert client.patch(
        f"/api/tasks/{completed.json['data']['id']}/status",
        json={"status": "completed"},
        headers={"X-CSRF-Token": token},
    ).status_code == 200
    metrics = _verified_metrics([
        {"status": "todo", "due_date": date.today() - timedelta(days=1)},
        {"status": "completed", "due_date": date.today() - timedelta(days=1)},
    ])
    assert metrics == {"total": 2, "completed": 1, "pending": 1, "overdue": 1, "completion_rate": 50}
    excel = client.get("/api/reports/tasks.xlsx?priority=high")
    assert excel.status_code == 200
    rows = list(load_workbook(BytesIO(excel.data), read_only=True)["Task Report"].iter_rows(min_row=5, values_only=True))
    assert len(rows) == 1
    assert rows[0][1] == "High report task"


def test_settings_profile_preferences_and_password(client):
    register(client)
    assert client.get("/api/auth/preferences").json["data"]["preferences"]["default_priority"] == "medium"
    token = csrf(client)
    updated = client.put(
        "/api/auth/profile",
        json={"name": "Updated User", "email": "updated@example.com"},
        headers={"X-CSRF-Token": token},
    )
    assert updated.status_code == 200
    assert client.get("/api/auth/me").json["data"]["user"]["name"] == "Updated User"
    token = csrf(client)
    preferences = client.put(
        "/api/auth/preferences",
        json={"deadline_reminders": False, "overdue_alerts": True, "default_view": "calendar", "default_priority": "high"},
        headers={"X-CSRF-Token": token},
    )
    assert preferences.status_code == 200
    assert client.get("/api/auth/preferences").json["data"]["preferences"]["default_view"] == "calendar"
    token = csrf(client)
    assert client.post(
        "/api/auth/password",
        json={"current_password": "secure-pass-123", "new_password": "new-secure-pass-456", "password_confirmation": "new-secure-pass-456"},
        headers={"X-CSRF-Token": token},
    ).status_code == 200
    token = csrf(client)
    assert client.post("/api/auth/login", json={"email": "updated@example.com", "password": "new-secure-pass-456"}, headers={"X-CSRF-Token": token}).status_code == 200


def test_settings_validation_and_csrf(client):
    register(client)
    assert client.put("/api/auth/profile", json={"name": "", "email": "bad"}).status_code == 403
    token = csrf(client)
    assert client.put("/api/auth/profile", json={"name": "", "email": "bad"}, headers={"X-CSRF-Token": token}).status_code == 400
    token = csrf(client)
    assert client.put("/api/auth/preferences", json={"deadline_reminders": True, "overdue_alerts": True, "default_view": "unknown", "default_priority": "medium"}, headers={"X-CSRF-Token": token}).status_code == 400


def test_member_attendance_lifecycle_and_csrf(client):
    register(client)
    assert client.get("/api/attendance/status").status_code == 200
    assert client.get("/api/attendance/status").json["data"]["record"] is None
    assert client.post("/api/attendance/check-in").status_code == 403
    token = csrf(client)
    checked_in = client.post("/api/attendance/check-in", headers={"X-CSRF-Token": token})
    assert checked_in.status_code == 201
    assert checked_in.json["data"]["record"]["status"] == "checked_in"
    token = csrf(client)
    assert client.post("/api/attendance/check-in", headers={"X-CSRF-Token": token}).status_code == 409
    token = csrf(client)
    checked_out = client.post("/api/attendance/check-out", headers={"X-CSRF-Token": token})
    assert checked_out.status_code == 200
    assert checked_out.json["data"]["record"]["status"] == "checked_out"
    token = csrf(client)
    assert client.post("/api/attendance/check-out", headers={"X-CSRF-Token": token}).status_code == 409
    assert len(client.get("/api/attendance/history").json["data"]["records"]) == 1


def test_attendance_requires_check_in_and_is_user_scoped(client):
    register(client)
    token = csrf(client)
    assert client.post("/api/attendance/check-out", headers={"X-CSRF-Token": token}).status_code == 409
    token = csrf(client)
    client.post("/api/auth/logout", headers={"X-CSRF-Token": token})
    register(client, "two@example.com")
    assert client.get("/api/attendance/history").json["data"]["records"] == []
    assert client.get("/api/admin/attendance").status_code == 403
    token = csrf(client)
    client.post("/api/auth/logout", headers={"X-CSRF-Token": token})
    assert client.get("/api/attendance/status").status_code == 401


def test_admin_attendance_dashboard_counts_and_database_uniqueness(client):
    register(client)
    connection = mysql.connector.connect(**Config.DB_CONFIG)
    cursor = connection.cursor()
    cursor.execute("UPDATE users SET role = 'admin' WHERE email = %s", ("one@example.com",))
    connection.commit()
    cursor.close()
    connection.close()
    token = csrf(client)
    client.post("/api/auth/logout", headers={"X-CSRF-Token": token})
    register(client, "member@example.com")
    token = csrf(client)
    client.post("/api/auth/logout", headers={"X-CSRF-Token": token})
    token = csrf(client)
    client.post("/api/auth/login", json={"email": "one@example.com", "password": "secure-pass-123"}, headers={"X-CSRF-Token": token})
    token = csrf(client)
    assert client.post("/api/attendance/check-in", headers={"X-CSRF-Token": token}).status_code == 201
    dashboard = client.get("/api/admin/attendance")
    assert dashboard.status_code == 200
    assert dashboard.json["data"]["counts"]["members"] == 1
    assert dashboard.json["data"]["counts"]["checked_in"] == 1
    assert dashboard.json["data"]["counts"]["currently_checked_in"] == 1
    connection = mysql.connector.connect(**Config.DB_CONFIG)
    cursor = connection.cursor()
    with pytest.raises(mysql.connector.IntegrityError):
        cursor.execute(
            "INSERT INTO attendance (user_id, attendance_date, check_in_at) VALUES (%s, %s, %s)",
            (1, date.today(), datetime.now(timezone.utc).replace(tzinfo=None)),
        )
    connection.rollback()
    cursor.close()
    connection.close()


def test_attendance_timezone_conversion_and_business_date(client):
    client.application.config["APP_TIMEZONE"] = "Asia/Kolkata"
    near_midnight = datetime(2026, 10, 26, 19, 0, tzinfo=timezone.utc)
    with client.application.app_context():
        assert business_date(near_midnight).isoformat() == "2026-10-27"
        assert display_time(datetime(2026, 10, 26, 18, 30)) == "2026-10-27T00:00:00+05:30"


def test_concurrent_check_in_is_serialized_by_unique_constraint(client):
    register(client)
    with client.session_transaction() as session:
        user_id = session["user_id"]
    attendance_day = date.today()
    barrier = threading.Barrier(2)
    results = []

    def insert_record():
        connection = mysql.connector.connect(**Config.DB_CONFIG)
        cursor = connection.cursor()
        try:
            barrier.wait(timeout=5)
            cursor.execute(
                "INSERT INTO attendance (user_id, attendance_date, check_in_at) VALUES (%s, %s, %s)",
                (user_id, attendance_day, datetime.now(timezone.utc).replace(tzinfo=None)),
            )
            connection.commit()
            results.append("inserted")
        except mysql.connector.IntegrityError:
            connection.rollback()
            results.append("duplicate")
        finally:
            cursor.close()
            connection.close()

    threads = [threading.Thread(target=insert_record) for _ in range(2)]
    for thread in threads:
        thread.start()
    for thread in threads:
        thread.join(timeout=10)
    assert sorted(results) == ["duplicate", "inserted"]
